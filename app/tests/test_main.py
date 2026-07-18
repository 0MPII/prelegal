import io
import re

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.main import app, get_catalog_entry, load_catalog

client = TestClient(app)
CATALOG = load_catalog()
CATALOG_FILENAMES = [entry["filename"] for entry in CATALOG]


# --- catalog / index ---


def test_index_lists_every_catalog_document_with_a_create_link():
    resp = client.get("/")
    assert resp.status_code == 200
    for entry in CATALOG:
        assert entry["name"] in resp.text
        assert f'/create/{entry["filename"]}' in resp.text


def test_get_catalog_entry_rejects_unknown_filenames_before_touching_disk():
    # This is the single gatekeeper every filesystem-touching route relies on;
    # if it doesn't reject unknown names, /template and /create could be
    # coaxed into reading arbitrary files via the filename path param.
    with pytest.raises(HTTPException) as exc_info:
        get_catalog_entry("../catalog.json")
    assert exc_info.value.status_code == 404


# --- template view ---


@pytest.mark.parametrize("filename", CATALOG_FILENAMES)
def test_template_view_renders_for_every_catalog_document(filename):
    resp = client.get(f"/template/{filename}")
    assert resp.status_code == 200
    assert "Fill in and create this document" in resp.text


def test_template_view_404s_for_unknown_document():
    resp = client.get("/template/does-not-exist.md")
    assert resp.status_code == 404


# --- create form ---


@pytest.mark.parametrize("filename", CATALOG_FILENAMES)
def test_create_form_renders_for_every_catalog_document(filename):
    resp = client.get(f"/create/{filename}")
    assert resp.status_code == 200
    assert "<form" in resp.text
    assert 'action="/create/' in resp.text


def test_create_form_404s_for_unknown_document():
    resp = client.get("/create/does-not-exist.md")
    assert resp.status_code == 404


# --- preview ---


@pytest.mark.parametrize("filename", CATALOG_FILENAMES)
def test_preview_fills_in_every_submitted_field_for_every_catalog_document(filename):
    form_resp = client.get(f"/create/{filename}")
    field_keys = re.findall(r'name="([a-z0-9_]+)"', form_resp.text)
    assert field_keys, "expected at least one input field on the create form"

    payload = {key: f"test-value-{key}" for key in field_keys}
    resp = client.post(f"/create/{filename}/preview", data=payload)
    assert resp.status_code == 200
    assert 'class="filled-field"' in resp.text
    assert 'class="empty-field"' not in resp.text  # every field was supplied


def test_preview_marks_unfilled_fields_as_bracketed_placeholders():
    resp = client.post("/create/Mutual-NDA.md/preview", data={})
    assert resp.status_code == 200
    assert 'class="empty-field"' in resp.text
    assert "[Purpose]" in resp.text


def test_preview_404s_for_unknown_document():
    resp = client.post("/create/does-not-exist.md/preview", data={})
    assert resp.status_code == 404


# --- download ---


@pytest.mark.parametrize("filename", CATALOG_FILENAMES)
def test_download_returns_a_valid_docx_for_every_catalog_document(filename):
    resp = client.post(f"/create/{filename}/download", data={})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment; filename=" in resp.headers["content-disposition"]

    doc = Document(io.BytesIO(resp.content))
    assert len(doc.tables) == 1


def test_download_filename_is_slugified_from_document_and_party_names():
    resp = client.post(
        "/create/Mutual-NDA.md/download",
        data={"party_1_name": "Acme Inc.", "party_2_name": "Beta LLC"},
    )
    assert resp.status_code == 200
    disposition = resp.headers["content-disposition"]
    assert "acme_inc" in disposition
    assert "beta_llc" in disposition


def test_download_falls_back_to_role_name_when_party_name_is_blank():
    resp = client.post("/create/Mutual-NDA.md/download", data={})
    assert resp.status_code == 200
    disposition = resp.headers["content-disposition"].lower()
    assert "party_1" in disposition
    assert "party_2" in disposition


def test_download_404s_for_unknown_document():
    resp = client.post("/create/does-not-exist.md/download", data={})
    assert resp.status_code == 404
